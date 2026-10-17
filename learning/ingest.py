"""
Document ingestion: upload -> extract text -> split into chunks -> index.

Supported: PDF, DOCX, PPTX, TXT, Markdown. Only the extracted TEXT is stored
(in the document_chunks table); the original file is not kept on disk, so
there is no uploaded file on the server that could be executed or fetched.

"Indexing" here means the chunks are stored with their page/slide and section
and are searched with TF-IDF at question time (see retrieval.py). No
embedding model or vector database is involved.

Processing runs on a small background thread pool so the upload request
returns immediately and the page can show progress by polling the document's
`stage` field.
"""

import io
import re
from concurrent.futures import ThreadPoolExecutor

import db
from learning import textproc

ALLOWED = {".pdf", ".docx", ".pptx", ".txt", ".md", ".markdown"}
MAX_BYTES = 15 * 1024 * 1024
MAX_PAGES = 600
CHUNK_TARGET = 750
CHUNK_MAX = 1300

_pool = ThreadPoolExecutor(max_workers=2, thread_name_prefix="ingest")
SYNC = False   # tests set this to True so processing finishes before the request returns


class IngestError(Exception):
    """Raised with a message that is safe to show to the student."""


def extension(filename: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", filename or "")
    return m.group(1).lower() if m else ""


# ------------------------------------------------------------ extraction
# Each extractor returns a list of units: {"page": int|None, "title": str|None, "text": str}

def _extract_pdf(data: bytes):
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise IngestError("This PDF is password-protected. Remove the password and upload it again.")
        pages = reader.pages
        if len(pages) > MAX_PAGES:
            raise IngestError(f"This PDF has more than {MAX_PAGES} pages. Please split it into smaller files.")
        return [{"page": i, "title": None, "text": (p.extract_text() or "")} for i, p in enumerate(pages, start=1)]
    except IngestError:
        raise
    except Exception:
        raise IngestError("I couldn't read this PDF. It may be damaged.")


def _extract_docx(data: bytes):
    import docx
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception:
        raise IngestError("I couldn't read this Word file. It may be damaged or not a .docx file.")
    lines = []
    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt:
            lines.append("")
        elif (p.style.name or "").lower().startswith(("heading", "title")):
            lines.append("\n" + txt + "\n")
        else:
            lines.append(txt)
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return [{"page": None, "title": None, "text": "\n".join(lines)}]


def _extract_pptx(data: bytes):
    from pptx import Presentation
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception:
        raise IngestError("I couldn't read this PowerPoint file. It may be damaged or not a .pptx file.")
    units = []
    for i, slide in enumerate(prs.slides, start=1):
        title = None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip() or None
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append("\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip()))
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Speaker notes: " + notes)
        units.append({"page": i, "title": title, "text": "\n\n".join(p for p in parts if p)})
    return units


def _extract_text(data: bytes):
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return [{"page": None, "title": None, "text": data.decode(enc)}]
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise IngestError("I couldn't read this text file.")


EXTRACTORS = {".pdf": _extract_pdf, ".docx": _extract_docx, ".pptx": _extract_pptx,
              ".txt": _extract_text, ".md": _extract_text, ".markdown": _extract_text}


# --------------------------------------------------------------- chunking

def _md_heading(line: str):
    m = re.match(r"^#{1,6}\s+(.+?)\s*#*$", line.strip())
    return m.group(1) if m else None


def chunk_units(units):
    """Turn extracted units into [{page, section, text}] chunks of roughly
    CHUNK_TARGET characters, never mixing two sections in one chunk."""
    chunks = []
    section = None
    for u in units:
        if u["title"]:
            section = u["title"]
        page = u["page"]
        buf, buf_section = [], section

        def flush():
            nonlocal buf
            text = "\n".join(buf).strip()
            if text:
                for piece in _split_long(text):
                    chunks.append({"page": page, "section": buf_section, "text": piece})
            buf = []

        lines = textproc.clean(u["text"]).split("\n")
        for i, raw in enumerate(lines):
            line = raw.strip()
            heading = _md_heading(line)
            if heading is None and line:
                prev_blank = i == 0 or not lines[i - 1].strip()
                next_blank = i + 1 >= len(lines) or not lines[i + 1].strip()
                single = len(line.split()) == 1 and line[0].isupper() and line.isalpha() and len(line) >= 4
                if textproc.is_heading(line) or (single and prev_blank and next_blank):
                    heading = line
            if heading:
                flush()
                section = heading.strip("# ").strip()
                buf_section = section
                continue
            buf.append(line)
            if sum(len(b) for b in buf) >= CHUNK_TARGET and not line:
                flush()
                buf_section = section
        flush()
    return chunks


def _split_long(text: str):
    if len(text) <= CHUNK_MAX:
        return [text]
    pieces, cur = [], ""
    for sent in re.split(r"(?<=[.!?])\s+", text):
        if cur and len(cur) + len(sent) > CHUNK_TARGET:
            pieces.append(cur.strip())
            cur = ""
        cur += sent + " "
    if cur.strip():
        pieces.append(cur.strip())
    return pieces


# ------------------------------------------------------------ the pipeline

def _stage(doc_id: int, stage: str) -> None:
    db.execute("UPDATE documents SET stage = ? WHERE id = ?", (stage, doc_id))


def process_document(doc_id: int, data: bytes, ext: str) -> None:
    """Runs on a background thread. Never raises: failures are recorded on
    the document row so the page can show them."""
    try:
        doc = db.fetchone("SELECT user_id, course_id FROM documents WHERE id = ?", (doc_id,))
        if doc is None:
            return
        _stage(doc_id, "Extracting text...")
        units = EXTRACTORS[ext](data)
        if not any(u["text"].strip() for u in units):
            raise IngestError("I couldn't find any text in this file. If it is a scan or photo, it needs OCR, which Nocturna doesn't have yet.")
        _stage(doc_id, "Creating chunks...")
        chunks = chunk_units(units)
        if not chunks:
            raise IngestError("This file doesn't contain enough readable text.")
        _stage(doc_id, "Indexing...")
        rows = [(doc_id, doc["user_id"], doc["course_id"], i, c["page"], c["section"], c["text"])
                for i, c in enumerate(chunks)]
        db.executemany(
            "INSERT INTO document_chunks (document_id, user_id, course_id, ord, page, section, text) VALUES (?,?,?,?,?,?,?)",
            rows,
        )
        db.execute(
            "UPDATE documents SET status = 'ready', stage = 'Ready', pages = ?, chunk_count = ?, error = NULL WHERE id = ?",
            (len([u for u in units if u["page"]]) or len(units), len(chunks), doc_id),
        )
    except IngestError as e:
        db.execute("UPDATE documents SET status = 'failed', stage = 'Failed', error = ? WHERE id = ?", (str(e), doc_id))
    except Exception:
        db.execute(
            "UPDATE documents SET status = 'failed', stage = 'Failed', error = ? WHERE id = ?",
            ("Something went wrong while processing this file. Please try again.", doc_id),
        )


def submit(user_id: int, filename: str, data: bytes, course_id=None, kind="material", wait=False) -> int:
    """Validate, create the document row, and start processing. Returns the id."""
    ext = extension(filename)
    if ext not in ALLOWED:
        raise IngestError("That file type isn't supported. Upload a PDF, DOCX, PPTX, TXT or Markdown file.")
    if not data:
        raise IngestError("That file is empty.")
    if len(data) > MAX_BYTES:
        raise IngestError(f"That file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    safe_name = re.sub(r"[^\w .()\-]", "_", filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1])[:120] or "document"
    doc_id = db.execute(
        "INSERT INTO documents (user_id, course_id, filename, kind, stage) VALUES (?, ?, ?, ?, 'Uploading...')",
        (user_id, course_id, safe_name, kind),
    )
    if wait or SYNC:
        process_document(doc_id, data, ext)
    else:
        _pool.submit(process_document, doc_id, data, ext)
    return doc_id
